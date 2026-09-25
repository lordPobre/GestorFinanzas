from decimal import Decimal
from io import BytesIO

from .models import Categoria, Transaccion

MESES = ['enero', 'febrero', 'marzo', 'abril', 'mayo', 'junio', 'julio',
         'agosto', 'septiembre', 'octubre', 'noviembre', 'diciembre']

ORIGENES = {
    'ingreso': 'Ingreso',
    'cuota': 'Cuota de compra',
    'suscripcion': 'Suscripción',
    'gasto': 'Gasto',
}

COLOR_ORIGEN = {
    'Ingreso': '1E7B32',
    'Cuota de compra': '4B4FC4',
    'Suscripción': 'B25E00',
    'Gasto': '4A4A4A',
}

VERDE = '1E7B32'
ROJO = 'B3261E'
TINTA = '2B2B2B'
AMBAR_TEXTO = '8A5200'
AMBAR_FONDO = 'FFF1DB'
AMBAR_LINEA = 'E8C88A'
GRIS_FONDO = 'F4F4F4'
GRIS_LINEA = 'D9D9D9'

FORMATO_MONTO = '"$"#,##0'
FORMATO_FECHA = 'DD/MM/YYYY'

MAX_FILAS = 20_000


def nombre_mes(anio, mes):
    return f'{MESES[mes - 1]} {anio}'.capitalize()


def filas_movimientos(usuario):
    etiquetas = dict(Categoria.opciones(usuario))
    filas = []

    consulta = Transaccion.objects.filter(usuario=usuario)
    total = consulta.count()
    if total > MAX_FILAS:
        movimientos = list(consulta.order_by('-fecha', '-id')[:MAX_FILAS])
        movimientos.reverse()
    else:
        movimientos = consulta.order_by('fecha', 'id').iterator(chunk_size=2000)

    for t in movimientos:
        categoria = etiquetas.get(t.categoria, t.categoria)
        es_ingreso = t.es_ingreso
        filas.append({
            'periodo': f'{t.fecha.year}-{t.fecha.month:02d}',
            'anio': t.fecha.year,
            'mes': t.fecha.month,
            'mes_nombre': nombre_mes(t.fecha.year, t.fecha.month),
            'fecha': t.fecha,
            'origen': ORIGENES.get(t.clase_tono, 'Gasto'),
            'tipo': t.get_tipo_display(),
            'categoria': categoria,
            'descripcion': t.descripcion or categoria,
            'estado': 'Recibido' if es_ingreso else ('Pagado' if t.pagado else 'Sin pagar'),
            'fecha_pago': t.fecha_pago,
            'ingreso': t.monto if es_ingreso else Decimal('0'),
            'egreso': Decimal('0') if es_ingreso else t.monto,
        })
    return filas


def resumen_por_mes(filas):
    meses = {}
    for f in filas:
        m = meses.setdefault(f['periodo'], {
            'periodo': f['periodo'], 'mes_nombre': f['mes_nombre'],
            'ingresos': Decimal('0'), 'egresos': Decimal('0'), 'cuenta': 0,
        })
        m['ingresos'] += f['ingreso']
        m['egresos'] += f['egreso']
        m['cuenta'] += 1
    return [meses[k] for k in sorted(meses)]


COLUMNAS_CSV = ['Periodo', 'Fecha', 'Corresponde a', 'Tipo', 'Categoria',
                'Descripcion', 'Estado', 'Fecha de pago',
                'Ingreso', 'Egreso', 'Balance']


def escribir_csv(writer, usuario, cuenta, hoy):
    filas = filas_movimientos(usuario)
    escribir = writer.writerow

    escribir(['Rekon — Movimientos'])
    escribir(['Cuenta', cuenta])
    escribir(['Exportado', hoy.strftime('%d/%m/%Y')])
    escribir(['Movimientos', len(filas)])
    if filas:
        escribir(['Desde', filas[0]['mes_nombre']])
        escribir(['Hasta', filas[-1]['mes_nombre']])
    escribir([])
    escribir(COLUMNAS_CSV)

    def total(etiqueta, ingreso, egreso):
        escribir(['', '', '', '', '', etiqueta, '', '',
                  int(ingreso), int(egreso), int(ingreso - egreso)])

    total_in = total_eg = Decimal('0')
    mes = None
    mes_nombre_actual = ''
    mes_in = mes_eg = Decimal('0')

    for f in filas:
        if f['periodo'] != mes:
            if mes is not None:
                total(f'Subtotal {mes_nombre_actual}', mes_in, mes_eg)
            mes = f['periodo']
            mes_nombre_actual = f['mes_nombre']
            mes_in = mes_eg = Decimal('0')
            escribir([])
            escribir([mes, f['mes_nombre'].upper()])

        mes_in += f['ingreso']
        mes_eg += f['egreso']
        total_in += f['ingreso']
        total_eg += f['egreso']

        escribir([
            f['periodo'],
            f['fecha'].strftime('%d/%m/%Y'),
            f['origen'],
            f['tipo'],
            f['categoria'],
            f['descripcion'],
            f['estado'],
            f['fecha_pago'].strftime('%d/%m/%Y') if f['fecha_pago'] else '',
            int(f['ingreso']) if f['ingreso'] else '',
            int(f['egreso']) if f['egreso'] else '',
            '',
        ])

    if mes is not None:
        total(f'Subtotal {mes_nombre_actual}', mes_in, mes_eg)
        escribir([])
        total('TOTAL GENERAL', total_in, total_eg)


COLUMNAS_HOJA = [
    ('Fecha', 12), ('Corresponde a', 17), ('Categoria', 26),
    ('Descripcion', 40), ('Estado', 13), ('Fecha de pago', 14),
    ('Ingreso', 14), ('Egreso', 14), ('Balance', 14),
]

COLUMNAS_DATOS = [
    ('Periodo', 10), ('Fecha', 12), ('Corresponde a', 17), ('Tipo', 10),
    ('Categoria', 26), ('Descripcion', 40), ('Estado', 13),
    ('Fecha de pago', 14), ('Ingreso', 14), ('Egreso', 14),
]


def libro_excel(usuario, cuenta, hoy):
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

    filas = filas_movimientos(usuario)
    resumen = resumen_por_mes(filas)

    fuente = 'Calibri'
    negrita_blanca = Font(name=fuente, size=10.5, bold=True, color='FFFFFF')
    linea_ambar = Border(top=Side(style='thin', color=AMBAR_LINEA))
    linea_gris = Border(top=Side(style='thin', color=GRIS_LINEA))
    fondo_tinta = PatternFill('solid', fgColor=TINTA)
    fondo_ambar = PatternFill('solid', fgColor=AMBAR_FONDO)
    fondo_gris = PatternFill('solid', fgColor=GRIS_FONDO)

    wb = Workbook()

    ws = wb.active
    ws.title = 'Movimientos'
    ws.sheet_view.showGridLines = False

    ws['A1'] = 'Rekon — Movimientos'
    ws['A1'].font = Font(name=fuente, size=15, bold=True, color=AMBAR_TEXTO)
    ws['A2'] = f'Cuenta: {cuenta}'
    ws['A3'] = f'Exportado el {hoy.strftime("%d/%m/%Y")} · {len(filas)} movimiento'\
               f'{"s" if len(filas) != 1 else ""}'
    if filas:
        ws['A4'] = f'{filas[0]["mes_nombre"]} — {filas[-1]["mes_nombre"]}'
    for celda in ('A2', 'A3', 'A4'):
        ws[celda].font = Font(name=fuente, size=10, color='6A6A6A')

    encabezado = 6
    for i, (titulo, ancho) in enumerate(COLUMNAS_HOJA, start=1):
        c = ws.cell(row=encabezado, column=i, value=titulo)
        c.font = negrita_blanca
        c.fill = fondo_tinta
        c.alignment = Alignment(horizontal='right' if i >= 7 else 'left',
                                vertical='center')
        ws.column_dimensions[c.column_letter].width = ancho
    ws.row_dimensions[encabezado].height = 22
    ws.freeze_panes = f'A{encabezado + 1}'

    fila = encabezado + 1
    total_in = total_eg = Decimal('0')

    def escribir_total(fila, etiqueta, ingreso, egreso, destacado=False):
        for col in range(1, len(COLUMNAS_HOJA) + 1):
            c = ws.cell(row=fila, column=col)
            c.fill = fondo_tinta if destacado else fondo_gris
            c.border = linea_gris
            c.font = negrita_blanca if destacado else Font(
                name=fuente, size=10.5, bold=True, color=TINTA)
        c = ws.cell(row=fila, column=4, value=etiqueta)
        c.alignment = Alignment(horizontal='right')
        for col, valor in ((7, ingreso), (8, egreso), (9, ingreso - egreso)):
            c = ws.cell(row=fila, column=col, value=int(valor))
            c.number_format = FORMATO_MONTO
            c.alignment = Alignment(horizontal='right')
        ws.row_dimensions[fila].height = 19

    mes = None
    mes_in = mes_eg = Decimal('0')
    mes_nombre_actual = ''

    for f in filas:
        if f['periodo'] != mes:
            if mes is not None:
                escribir_total(fila, f'Subtotal · {mes_nombre_actual}', mes_in, mes_eg)
                fila += 2
            mes = f['periodo']
            mes_nombre_actual = f['mes_nombre']
            mes_in = mes_eg = Decimal('0')

            for col in range(1, len(COLUMNAS_HOJA) + 1):
                c = ws.cell(row=fila, column=col)
                c.fill = fondo_ambar
                c.border = linea_ambar
            c = ws.cell(row=fila, column=1, value=f['mes_nombre'].upper())
            c.font = Font(name=fuente, size=11, bold=True, color=AMBAR_TEXTO)
            c.alignment = Alignment(vertical='center')
            ws.row_dimensions[fila].height = 24
            fila += 1

        mes_in += f['ingreso']
        mes_eg += f['egreso']
        total_in += f['ingreso']
        total_eg += f['egreso']

        normal = Font(name=fuente, size=10.5, color=TINTA)

        c = ws.cell(row=fila, column=1, value=f['fecha'])
        c.number_format = FORMATO_FECHA
        c.font = normal

        c = ws.cell(row=fila, column=2, value=f['origen'])
        c.font = Font(name=fuente, size=10.5, bold=True,
                      color=COLOR_ORIGEN.get(f['origen'], TINTA))

        ws.cell(row=fila, column=3, value=f['categoria']).font = normal
        ws.cell(row=fila, column=4, value=f['descripcion']).font = normal

        c = ws.cell(row=fila, column=5, value=f['estado'])
        c.font = Font(name=fuente, size=10.5, bold=f['estado'] == 'Sin pagar',
                      color=ROJO if f['estado'] == 'Sin pagar' else VERDE)

        if f['fecha_pago']:
            c = ws.cell(row=fila, column=6, value=f['fecha_pago'])
            c.number_format = FORMATO_FECHA
            c.font = Font(name=fuente, size=10.5, color='6A6A6A')

        if f['ingreso']:
            c = ws.cell(row=fila, column=7, value=int(f['ingreso']))
            c.number_format = FORMATO_MONTO
            c.font = Font(name=fuente, size=10.5, bold=True, color=VERDE)
        if f['egreso']:
            c = ws.cell(row=fila, column=8, value=int(f['egreso']))
            c.number_format = FORMATO_MONTO
            c.font = Font(name=fuente, size=10.5, color=ROJO)

        fila += 1

    if mes is not None:
        escribir_total(fila, f'Subtotal · {mes_nombre_actual}', mes_in, mes_eg)
        escribir_total(fila + 2, 'TOTAL GENERAL', total_in, total_eg, destacado=True)

    hoja = wb.create_sheet('Resumen por mes')
    hoja.sheet_view.showGridLines = False
    for i, (titulo, ancho) in enumerate(
            [('Periodo', 10), ('Mes', 22), ('Ingresos', 14),
             ('Egresos', 14), ('Balance', 14), ('Movimientos', 13)], start=1):
        c = hoja.cell(row=1, column=i, value=titulo)
        c.font = negrita_blanca
        c.fill = fondo_tinta
        c.alignment = Alignment(horizontal='right' if i >= 3 else 'left',
                                vertical='center')
        hoja.column_dimensions[c.column_letter].width = ancho
    hoja.row_dimensions[1].height = 22
    hoja.freeze_panes = 'A2'

    for n, m in enumerate(resumen, start=2):
        balance = m['ingresos'] - m['egresos']
        hoja.cell(row=n, column=1, value=m['periodo'])
        hoja.cell(row=n, column=2, value=m['mes_nombre'])
        for col, valor, color in ((3, m['ingresos'], VERDE),
                                  (4, m['egresos'], ROJO),
                                  (5, balance, VERDE if balance >= 0 else ROJO)):
            c = hoja.cell(row=n, column=col, value=int(valor))
            c.number_format = FORMATO_MONTO
            c.font = Font(name=fuente, size=10.5, bold=col == 5, color=color)
        hoja.cell(row=n, column=6, value=m['cuenta']).alignment = Alignment(
            horizontal='right')

    plana = wb.create_sheet('Datos')
    for i, (titulo, ancho) in enumerate(COLUMNAS_DATOS, start=1):
        c = plana.cell(row=1, column=i, value=titulo)
        c.font = negrita_blanca
        c.fill = fondo_tinta
        plana.column_dimensions[c.column_letter].width = ancho
    plana.freeze_panes = 'A2'

    for n, f in enumerate(filas, start=2):
        plana.cell(row=n, column=1, value=f['periodo'])
        c = plana.cell(row=n, column=2, value=f['fecha'])
        c.number_format = FORMATO_FECHA
        plana.cell(row=n, column=3, value=f['origen'])
        plana.cell(row=n, column=4, value=f['tipo'])
        plana.cell(row=n, column=5, value=f['categoria'])
        plana.cell(row=n, column=6, value=f['descripcion'])
        plana.cell(row=n, column=7, value=f['estado'])
        if f['fecha_pago']:
            c = plana.cell(row=n, column=8, value=f['fecha_pago'])
            c.number_format = FORMATO_FECHA
        for col, valor in ((9, f['ingreso']), (10, f['egreso'])):
            c = plana.cell(row=n, column=col, value=int(valor))
            c.number_format = FORMATO_MONTO
    if filas:
        plana.auto_filter.ref = f'A1:J{len(filas) + 1}'

    buffer = BytesIO()
    wb.save(buffer)
    return buffer.getvalue()

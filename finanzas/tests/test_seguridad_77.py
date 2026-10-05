import io
import zipfile
from datetime import date
from unittest import mock

from django.contrib.auth.models import User
from django.contrib.messages.storage.fallback import FallbackStorage
from django.contrib.sessions.middleware import SessionMiddleware
from django.test import RequestFactory, SimpleTestCase, TestCase
from django.urls import reverse

from core.settings import limpiar_evento_sentry
from core.urls import entrada_admin, permiso_admin

from ..cartolas.base import ErrorCartola
from ..cartolas.tabla import _revisar_zip
from ..exportar import libro_excel, texto_csv
from ..models import SegundoFactor, Transaccion
from ..redirecciones import destino_seguro


class DestinoSeguroTests(SimpleTestCase):

    def setUp(self):
        self.pedido = RequestFactory().get('/', HTTP_HOST='testserver')

    def test_acepta_rutas_propias(self):
        for valor in ('/metas/', '/cuotas/?mes=3', '/perfil/#seguridad'):
            self.assertEqual(destino_seguro(self.pedido, valor), valor)

    def test_rechaza_destinos_externos(self):
        for valor in ('//malo.com', '/\\malo.com', '\\\\malo.com', 'https://malo.com',
                      'javascript:alert(1)', '/\tmalo', ' //malo.com', '', None):
            self.assertEqual(destino_seguro(self.pedido, valor, 'X'), 'X')


class RedireccionesTests(TestCase):

    def setUp(self):
        self.usuario = User.objects.create_user('ana', password='clave-larga-123')

    def test_login_no_sigue_destinos_externos(self):
        r = self.client.post(reverse('login'), {
            'username': 'ana', 'password': 'clave-larga-123', 'next': '/\\malo.com'})
        self.assertEqual(r['Location'], reverse('dashboard'))

    def test_login_sigue_destino_propio(self):
        r = self.client.post(reverse('login'), {
            'username': 'ana', 'password': 'clave-larga-123', 'next': '/metas/'})
        self.assertEqual(r['Location'], '/metas/')

    def test_next_path_externo_se_ignora(self):
        self.client.force_login(self.usuario)
        r = self.client.post(reverse('crear_gasto_pendiente'), {
            'next': '?mes=1', 'next_path': '//malo.com/'})
        self.assertFalse(r['Location'].startswith('//'))
        self.assertTrue(r['Location'].endswith('?mes=1'))


class ExportacionTests(TestCase):

    def test_csv_neutraliza_formulas(self):
        self.assertEqual(texto_csv('=HYPERLINK("x")'), '\'=HYPERLINK("x")')
        self.assertEqual(texto_csv('+56 9'), "'+56 9")
        self.assertEqual(texto_csv('@SUMA'), "'@SUMA")
        self.assertEqual(texto_csv('Supermercado'), 'Supermercado')
        self.assertEqual(texto_csv(1500), 1500)

    def test_excel_guarda_formulas_como_texto(self):
        from openpyxl import load_workbook

        usuario = User.objects.create_user('beto', password='x')
        Transaccion.objects.create(usuario=usuario, tipo='EGRESO', monto=1000,
                                   categoria='Otros', descripcion='=1+1',
                                   fecha=date(2026, 9, 1))
        libro = load_workbook(io.BytesIO(libro_excel(usuario, 'beto', date(2026, 10, 1))))
        celdas = [c for hoja in libro for fila in hoja.iter_rows() for c in fila
                  if c.value == '=1+1']
        self.assertTrue(celdas)
        self.assertTrue(all(c.data_type == 's' for c in celdas))


class ZipExcelTests(SimpleTestCase):

    def _zip(self, tamano):
        crudo = io.BytesIO()
        with zipfile.ZipFile(crudo, 'w', zipfile.ZIP_DEFLATED) as z:
            z.writestr('xl/worksheets/sheet1.xml', b'0' * tamano)
        crudo.seek(0)
        return crudo

    def test_rechaza_zip_bomba(self):
        with self.assertRaises(ErrorCartola):
            _revisar_zip(self._zip(60 * 1024 * 1024))

    def test_acepta_excel_normal(self):
        binario = self._zip(200 * 1024)
        _revisar_zip(binario)
        self.assertEqual(binario.tell(), 0)

    def test_rechaza_lo_que_no_es_zip(self):
        with self.assertRaises(ErrorCartola):
            _revisar_zip(io.BytesIO(b'no soy un excel'))


class SinCacheTests(TestCase):

    def test_paginas_con_sesion_no_se_guardan(self):
        self.client.force_login(User.objects.create_user('caro', password='x'))
        r = self.client.get(reverse('perfil'))
        self.assertIn('no-store', r['Cache-Control'])
        self.assertIn('private', r['Cache-Control'])

    def test_paginas_publicas_no_cambian(self):
        r = self.client.get(reverse('privacidad'))
        self.assertNotIn('no-store', r.get('Cache-Control', ''))


class SentryTests(SimpleTestCase):

    def test_quita_datos_del_pedido(self):
        evento = {
            'request': {
                'url': 'https://fintora.cl/recuperar/MQ/abc-123/',
                'data': {'monto': '50000'},
                'cookies': {'sessionid': 'x'},
                'query_string': 'ahorro=1',
                'headers': {'Cookie': 'x', 'User-Agent': 'Safari', 'Authorization': 'y'},
            },
            'user': {'id': 3},
        }
        limpio = limpiar_evento_sentry(evento, {})
        pedido = limpio['request']
        self.assertNotIn('data', pedido)
        self.assertNotIn('cookies', pedido)
        self.assertNotIn('query_string', pedido)
        self.assertEqual(pedido['headers'], {'User-Agent': 'Safari'})
        self.assertNotIn('abc-123', pedido['url'])
        self.assertNotIn('user', limpio)


class RutasIaTests(TestCase):

    def setUp(self):
        self.client.force_login(User.objects.create_user('dani', password='x'))

    def test_get_no_permitido(self):
        self.assertEqual(self.client.get(reverse('analisis_ia')).status_code, 405)
        self.assertEqual(self.client.get(reverse('plan_ia')).status_code, 405)

    def test_post_sin_csrf_rechazado(self):
        from django.test import Client
        cliente = Client(enforce_csrf_checks=True)
        cliente.force_login(User.objects.get(username='dani'))
        self.assertEqual(cliente.post(reverse('analisis_ia')).status_code, 403)


class Admin2faTests(TestCase):

    def _pedido(self, usuario):
        pedido = RequestFactory().get('/panel/login/')
        SessionMiddleware(lambda r: None).process_request(pedido)
        pedido.session.save()
        pedido._messages = FallbackStorage(pedido)
        pedido.user = usuario
        return pedido

    def setUp(self):
        self.staff = User.objects.create_user('jefe', password='x', is_staff=True)
        parche = mock.patch('core.urls.reverse', return_value='/panel/')
        parche.start()
        self.addCleanup(parche.stop)

    def test_staff_sin_2fa_va_a_activarla(self):
        r = entrada_admin(self._pedido(self.staff))
        self.assertEqual(r['Location'], reverse('configurar_2fa'))
        self.assertFalse(permiso_admin(self._pedido(self.staff)))

    def test_staff_con_2fa_entra(self):
        SegundoFactor.objects.create(usuario=self.staff, secreto=SegundoFactor.generar_secreto(),
                                     activo=True)
        r = entrada_admin(self._pedido(self.staff))
        self.assertEqual(r['Location'], '/panel/')
        self.assertTrue(permiso_admin(self._pedido(self.staff)))

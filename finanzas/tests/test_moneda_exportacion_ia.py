import csv
import os
from datetime import date
from decimal import Decimal
from io import BytesIO, StringIO
from types import SimpleNamespace
from unittest import mock

from django.contrib.auth.models import User
from django.test import SimpleTestCase, TestCase

from .. import ia
from ..analisis import analizar_finanzas
from ..exportar import escribir_csv, filas_movimientos, libro_excel, nombre_mes, resumen_por_mes
from ..models import Transaccion
from ..templatetags.moneda import a_json, money, money_corto, money_signed, pct


class FiltrosDeMoneda(SimpleTestCase):

    def test_money_sin_decimales(self):
        self.assertEqual(money(1234567), '$1.234.567')
        self.assertEqual(money(-500), '-$500')
        self.assertEqual(money('no es numero'), '$0')

    def test_money_con_centavos(self):
        self.assertEqual(money(Decimal('12.45'), 'USD'), 'USD12,45')

    def test_money_signed(self):
        self.assertEqual(money_signed(500), '+$500')
        self.assertEqual(money_signed(-1000), '-$1.000')
        self.assertEqual(money_signed(None), '$0')

    def test_money_corto(self):
        self.assertEqual(money_corto(1234567), '$1,2M')
        self.assertEqual(money_corto(-2000000), '-$2M')
        self.assertEqual(money_corto(45000), '$45k')
        self.assertEqual(money_corto(999), '$999')
        self.assertEqual(money_corto('x'), '$0')

    def test_pct(self):
        self.assertEqual(pct(50, 200), 25)
        self.assertEqual(pct(300, 100), 100)
        self.assertEqual(pct(1, 0), 0)
        self.assertEqual(pct('a', 1), 0)

    def test_a_json(self):
        self.assertEqual(a_json({'a': Decimal('1.5')}), '{"a": 1.5}')
        self.assertEqual(a_json('[1, 2]'), '[1, 2]')
        self.assertEqual(a_json({'f': date(2026, 1, 2)}), '{"f": "2026-01-02"}')


class Exportacion(TestCase):

    def setUp(self):
        self.ana = User.objects.create_user('ana', 'ana@ejemplo.cl', 'clave-larga-1')
        Transaccion.objects.create(usuario=self.ana, tipo='INGRESO', monto=Decimal('1000000'),
                                   categoria='Sueldo', fecha=date(2026, 7, 5))
        Transaccion.objects.create(usuario=self.ana, tipo='EGRESO', monto=Decimal('25000'),
                                   categoria='Comida', descripcion='Almuerzo',
                                   fecha=date(2026, 7, 10), pagado=True,
                                   fecha_pago=date(2026, 7, 10))
        Transaccion.objects.create(usuario=self.ana, tipo='EGRESO', monto=Decimal('40000'),
                                   categoria='Comida', fecha=date(2026, 8, 3), pagado=False)

    def test_nombre_mes(self):
        self.assertEqual(nombre_mes(2026, 7), 'Julio 2026')

    def test_filas_y_resumen(self):
        filas = filas_movimientos(self.ana)
        self.assertEqual(len(filas), 3)
        self.assertEqual(filas[0]['ingreso'], Decimal('1000000'))
        self.assertEqual(filas[2]['estado'], 'Sin pagar')
        resumen = resumen_por_mes(filas)
        self.assertEqual([m['periodo'] for m in resumen], ['2026-07', '2026-08'])
        self.assertEqual(resumen[0]['egresos'], Decimal('25000'))
        self.assertEqual(resumen[1]['cuenta'], 1)

    def test_csv(self):
        salida = StringIO()
        escribir_csv(csv.writer(salida), self.ana, 'ana@ejemplo.cl', date(2026, 9, 1))
        texto = salida.getvalue()
        self.assertIn('Subtotal Julio 2026', texto)
        self.assertIn('Subtotal Agosto 2026', texto)
        self.assertIn('TOTAL GENERAL', texto)
        self.assertIn('Almuerzo', texto)
        self.assertIn('935000', texto)

    def test_csv_sin_movimientos(self):
        bea = User.objects.create_user('bea', 'bea@ejemplo.cl', 'clave-larga-2')
        salida = StringIO()
        escribir_csv(csv.writer(salida), bea, 'bea@ejemplo.cl', date(2026, 9, 1))
        self.assertNotIn('TOTAL GENERAL', salida.getvalue())

    def test_excel(self):
        from openpyxl import load_workbook

        libro = load_workbook(BytesIO(libro_excel(self.ana, 'ana@ejemplo.cl', date(2026, 9, 1))))
        self.assertEqual(libro.sheetnames, ['Movimientos', 'Resumen por mes', 'Datos'])
        self.assertEqual(libro['Datos'].max_row, 4)
        self.assertEqual(libro['Resumen por mes'].max_row, 3)
        self.assertEqual(libro['Datos'].auto_filter.ref, 'A1:J4')

    def test_excel_sin_movimientos(self):
        from openpyxl import load_workbook

        bea = User.objects.create_user('bea', 'bea@ejemplo.cl', 'clave-larga-2')
        libro = load_workbook(BytesIO(libro_excel(bea, 'bea@ejemplo.cl', date(2026, 9, 1))))
        self.assertEqual(libro['Datos'].max_row, 1)


ANALISIS = {
    'tiene_datos': True,
    'ingreso_mensual': 1000000,
    'gasto_mensual': 600000,
    'cuota_mensual_total': 150000,
    'deuda_total_restante': 900000,
    'flujo_libre': 250000,
    'dti': 15,
    'meses_restantes': 6,
    'cantidad_deudas': 2,
    'riesgo_nivel': 'bajo',
    'riesgo_score': 20,
    'tendencia': 'bajando',
    'riesgo_factores': [{'factor': 'DTI', 'detalle': 'bajo el 30%'}],
}


def _respuesta(texto):
    return SimpleNamespace(content=[SimpleNamespace(type='text', text=texto)])


class InterpretacionConIA(SimpleTestCase):

    def test_prompt_lleva_los_numeros(self):
        prompt = ia._construir_prompt(ANALISIS)
        self.assertIn('$1,000,000', prompt)
        self.assertIn('- DTI: bajo el 30%', prompt)

    def test_sin_clave_no_llama(self):
        with mock.patch.dict(os.environ, {'ANTHROPIC_API_KEY': ''}):
            self.assertIsNone(ia.interpretar_con_ia(ANALISIS))

    def test_sin_datos_no_llama(self):
        with mock.patch.dict(os.environ, {'ANTHROPIC_API_KEY': 'k'}):
            self.assertIsNone(ia.interpretar_con_ia({'tiene_datos': False}))

    def _llamar(self, respuesta=None, error=None):
        with mock.patch.dict(os.environ, {'ANTHROPIC_API_KEY': 'k'}), \
                mock.patch('anthropic.Anthropic') as cliente:
            crear = cliente.return_value.messages.create
            if error:
                crear.side_effect = error
            else:
                crear.return_value = respuesta
            return ia.interpretar_con_ia(ANALISIS)

    def test_respuesta_valida_en_bloque_de_codigo(self):
        datos = self._llamar(_respuesta('```json\n{"diagnostico": "ok", "recomendaciones": []}\n```'))
        self.assertEqual(datos['diagnostico'], 'ok')

    def test_respuesta_sin_claves(self):
        self.assertIsNone(self._llamar(_respuesta('{"otra": 1}')))

    def test_respuesta_que_no_es_json(self):
        self.assertIsNone(self._llamar(_respuesta('hola')))

    def test_falla_la_api(self):
        self.assertIsNone(self._llamar(error=RuntimeError('caida')))


class AnalisisFinanciero(TestCase):

    def test_sin_movimientos(self):
        ana = User.objects.create_user('ana', 'ana@ejemplo.cl', 'clave-larga-1')
        self.assertIsInstance(analizar_finanzas(ana), dict)

    def test_con_movimientos(self):
        ana = User.objects.create_user('ana', 'ana@ejemplo.cl', 'clave-larga-1')
        hoy = date.today()
        for meses in range(4):
            anio = hoy.year if hoy.month > meses else hoy.year - 1
            mes = (hoy.month - meses - 1) % 12 + 1
            fecha = date(anio, mes, 1)
            Transaccion.objects.create(usuario=ana, tipo='INGRESO', monto=Decimal('1000000'),
                                       categoria='Sueldo', fecha=fecha)
            Transaccion.objects.create(usuario=ana, tipo='EGRESO', monto=Decimal('300000'),
                                       categoria='Comida', fecha=fecha)
        resultado = analizar_finanzas(ana)
        self.assertIsInstance(resultado, dict)
        if resultado.get('tiene_datos'):
            self.assertIn('DATOS FINANCIEROS', ia._construir_prompt(resultado))
